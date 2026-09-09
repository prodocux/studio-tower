import io
from types import SimpleNamespace

import pytest
from app.models.action_proposal import (
    ActionExecutionRecord,
    ActionExecutionStatus,
    ActionSourceDescriptor,
    DeliverableFormat,
    issue_action_token,
    verify_action_token,
)
from app.models.deliverable_spec import DeliverableSpecError, resolve_deliverable_spec
from app.models.user import User
from app.services.deliverable_service import DeliverableExecutionService
from app.services.storage import store
from docx import Document
from pptx import Presentation
from pypdf import PdfReader


def test_output_format_is_covered_by_v3_action_signature():
    proposal = issue_action_token(
        space_id="space_format_signature",
        project_tag="general",
        user_id="format_user",
        title="Signed format",
        description="Explicit format contract",
        sources=[],
        action_type="create_scene_breakdown",
        output_format="pdf",
    )
    assert proposal.key_id == "v3"
    assert verify_action_token(proposal, "format_user", "space_format_signature") is True
    proposal.output_format = DeliverableFormat.CSV
    assert verify_action_token(proposal, "format_user", "space_format_signature") is False


def test_scene_breakdown_pdf_uses_readable_scene_cards_not_wide_pipe_rows():
    from app.services.prodocux_deliverable_renderer import render_deliverable

    user = User(uid="layout_user", email="layout@example.com", display_name="Layout Producer")
    proposal = issue_action_token(
        space_id="space_layout",
        project_tag="general",
        user_id=user.uid,
        title="Act III Scene Breakdown",
        description="Readable production breakdown",
        sources=[],
        action_type="create_scene_breakdown",
        output_format="pdf",
    )
    extracted = {
        "scenes": [{
            "scene_number": 1,
            "slugline": "EXT. HIGH ATMOSPHERE - STRATOSPHERE JUMP - NIGHT",
            "day_night": "NIGHT",
            "location": "HIGH ATMOSPHERE - STRATOSPHERE JUMP",
        }],
        "characters": ["MAYA", "REN"],
        "provenance": [{"file_id": "file_source", "generation": 1, "sha256": "abc", "chunk_count": 1}],
        "is_ungrounded": False,
    }

    _, _, payload, _, _ = render_deliverable(proposal, user, extracted)
    text = "\n".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(payload)).pages)
    assert "Scene 1" in text or "SCENE INDEX" in text
    assert "EXT. HIGH ATMOSPHERE" in text
    assert "MAYA" in text and "REN" in text
    assert "Scene | Slugline | Day/Night" not in text
    assert "Breakdown field" not in text


def test_character_extraction_does_not_promote_metadata_and_keywords(monkeypatch):
    source_text = """SCENE 1: EXT. ROOFTOP - NIGHT
PROJECT NEBULA RUNNER
PRODUCTION CONTROL
MAYA
We see Maya prepare for the jump.
REN
Ren checks the safety line.
TECHNICAL CAMERA DIRECTIVES
"""
    chunk = SimpleNamespace(normalized_text=source_text)
    monkeypatch.setattr(store, "get_document_chunks", lambda *_args, **_kwargs: [chunk])
    proposal = issue_action_token(
        space_id="space_character_scope",
        project_tag="general",
        user_id="character_user",
        title="Scoped characters",
        description="Do not confuse metadata with cast",
        sources=[ActionSourceDescriptor(
            file_id="file_script",
            active_generation=1,
            content_hash="hash",
            space_id="space_character_scope",
        )],
        action_type="create_scene_breakdown",
        output_format="json",
    )

    extracted = DeliverableExecutionService._extract_source_document_data(proposal)
    assert extracted["characters"] == ["MAYA", "REN"]


def test_scene_parser_separates_stage_notes_and_rejects_technical_terms_as_cast(monkeypatch):
    source_text = """SCENE 1: EXT. HIGH ATMOSPHERE - STRATOSPHERE JUMP - NIGHT | STAGE A
(LED VOLUME) & AERIAL SPLINTER UNIT
LEAD SHOWRUNNER
PRIMARY AI AGENTS
DETERMINISTIC GATES
OBSERVABILITY LINK
SECURITY DOMAIN
MAYA
Maya checks the tether before the jump.
SCENE 3: INT. TETHER SKYPORT - PLATFORM 42 - NIGHT | STAGE C (WIRE GANTRY & INDUSTRIAL RIG)
REN
Ren secures the wire harness.
"""
    chunk = SimpleNamespace(normalized_text=source_text)
    monkeypatch.setattr(store, "get_document_chunks", lambda *_args, **_kwargs: [chunk])
    proposal = issue_action_token(
        space_id="space_scene_semantics",
        project_tag="general",
        user_id="scene_user",
        title="Act III Breakdown",
        description="Separate screenplay fields from production notes",
        sources=[ActionSourceDescriptor(
            file_id="file_script",
            active_generation=1,
            content_hash="hash",
            space_id="space_scene_semantics",
        )],
        action_type="create_scene_breakdown",
        output_format="pdf",
    )

    extracted = DeliverableExecutionService._extract_source_document_data(proposal)
    assert extracted["characters"] == ["MAYA", "REN"]
    assert [scene["scene_number"] for scene in extracted["scenes"]] == ["1", "3"]
    assert extracted["scenes"][0]["slugline"] == "EXT. HIGH ATMOSPHERE - STRATOSPHERE JUMP - NIGHT"
    assert extracted["scenes"][0]["location"] == "HIGH ATMOSPHERE - STRATOSPHERE JUMP"
    assert "STAGE" not in extracted["scenes"][0]["slugline"]
    assert "Stunts / Safety" in extracted["production_elements"]
    assert "VFX / Virtual Production" in extracted["production_elements"]
    assert "Special Equipment" in extracted["production_elements"]
    assert extracted["scenes"][0]["characters"] == ["MAYA"]
    assert extracted["scenes"][1]["characters"] == ["REN"]
    assert "VFX / Virtual Production" in extracted["scenes"][0]["production_elements"]
    assert "Special Equipment" in extracted["scenes"][1]["production_elements"]


def test_exported_file_uses_gemini_decided_scenes(monkeypatch):
    from app.agent.brain import AgentBrain
    from app.agent.schemas import DeliverableContent, DeliverableScene
    from pypdf import PdfReader

    captured = {}

    def _compose(cls, **kwargs):
        captured.update(kwargs)
        return DeliverableContent(
            title="Act III Scene Breakdown",
            summary="Chronos Spire conflagration",
            scenes=[
                DeliverableScene(
                    scene_number="26",
                    slugline="INT. CONTROL APEX - HELIPAD ACCESS - DAWN",
                    day_night="DAWN",
                    location="CONTROL APEX - HELIPAD ACCESS",
                    characters=["MAYA", "KAELEN 9"],
                ),
                DeliverableScene(
                    scene_number="29",
                    slugline="EXT. CHRONOS SPIRE ROOFTOP - SHOWDOWN - DAWN",
                    day_night="DAWN",
                    location="CHRONOS SPIRE ROOFTOP - SHOWDOWN",
                    characters=["SILAS CROSS"],
                ),
            ],
            characters=["MAYA", "KAELEN 9", "SILAS CROSS"],
        )

    monkeypatch.setattr(AgentBrain, "compose_deliverable_content", classmethod(_compose))
    user = User(uid="gemini_file_user", email="ad@example.com", display_name="1st AD")
    proposal = issue_action_token(
        space_id="space_gemini_file",
        project_tag="general",
        user_id=user.uid,
        title="Act III Scene Breakdown",
        description="Only Act III",
        sources=[],
        action_type="create_scene_breakdown",
        output_format="pdf",
        metadata={"request_text": "Please make an Act III breakdown"},
    )
    filename, media_type, payload, _, _ = DeliverableExecutionService._generate_content_from_sources(proposal, user)
    text = "\n".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(payload)).pages)
    assert captured["action_type"] == "create_scene_breakdown"
    assert "Act III" in captured["user_request"]
    assert filename.endswith(".pdf")
    assert media_type == "application/pdf"
    assert "26" in text and "INT. CONTROL APEX" in text
    assert "29" in text and "CHRONOS SPIRE ROOFTOP" in text
    assert "HIGH ATMOSPHERE" not in text


def test_budget_uses_evidence_derived_accounts_without_inventing_prices():
    from app.services.prodocux_deliverable_renderer import render_deliverable

    user = User(uid="budget_evidence_user", email="budget-evidence@example.com")
    proposal = issue_action_token(
        space_id="space_budget_evidence",
        project_tag="general",
        user_id=user.uid,
        title="Act III Budget",
        description="Evidence-derived budget accounts",
        sources=[],
        action_type="export_production_budget",
        output_format="json",
    )
    extracted = {
        "scenes": [{"location": "TETHER SKYPORT"}],
        "characters": ["MAYA", "REN"],
        "production_elements": ["Stunts / Safety", "VFX / Virtual Production"],
        "is_ungrounded": False,
    }
    payload = render_deliverable(proposal, user, extracted)[2].decode("utf-8")
    assert "Cast: MAYA, REN" in payload
    assert "Locations: TETHER SKYPORT" in payload
    assert "Stunts / Safety" in payload
    assert "VFX / Virtual Production" in payload
    assert '"TBD"' in payload
    assert "15000" not in payload


def test_user_selected_word_budget_is_a_real_editable_table():
    from app.services.prodocux_deliverable_renderer import render_deliverable

    user = User(uid="budget_word_user", email="budget@example.com", display_name="Line Producer")
    proposal = issue_action_token(
        space_id="space_budget_word",
        project_tag="general",
        user_id=user.uid,
        title="Production Budget",
        description="Editable Word budget",
        sources=[],
        action_type="export_production_budget",
        output_format="docx",
    )
    filename, _, payload, _, _ = render_deliverable(proposal, user, {"is_ungrounded": True})
    document = Document(io.BytesIO(payload))
    assert filename.endswith(".docx")
    assert any("Account" in cell.text for table in document.tables for row in table.rows for cell in row.cells)


def test_user_selected_pptx_shot_list_uses_one_slide_per_shot():
    from app.services.prodocux_deliverable_renderer import render_deliverable

    user = User(uid="shot_pptx_user", email="shots@example.com", display_name="Director")
    proposal = issue_action_token(
        space_id="space_shot_pptx",
        project_tag="general",
        user_id=user.uid,
        title="Act I Shot List",
        description="Editable PowerPoint shot list",
        sources=[],
        action_type="generate_shot_list",
        output_format="pptx",
    )
    extracted = {
        "scenes": [
            {"scene_number": 1, "slugline": "EXT. ROOFTOP - NIGHT"},
            {"scene_number": 2, "slugline": "INT. CONTROL ROOM - NIGHT"},
        ],
        "is_ungrounded": False,
    }
    filename, _, payload, _, _ = render_deliverable(proposal, user, extracted)
    presentation = Presentation(io.BytesIO(payload))
    all_text = [shape.text for slide in presentation.slides for shape in slide.shapes if hasattr(shape, "text_frame")]
    assert filename.endswith(".pptx")
    assert any("Shot 1A" in text for text in all_text)
    assert any("Shot 2A" in text for text in all_text)
    assert not any("Shot # | Scene | Description" in text for text in all_text)


def test_dynamic_deliverable_spec_preserves_signed_user_presentation_intent():
    proposal = issue_action_token(
        space_id="space_dynamic_spec",
        project_tag="general",
        user_id="dynamic_user",
        title="Producer Working Budget",
        description="A detailed editable department budget in Word",
        sources=[],
        action_type="export_production_budget",
        output_format="docx",
        metadata={
            "deliverable_spec": {
                "audience": "producer",
                "purpose": "working_document",
                "layout": "department_tables",
                "density": "detailed",
                "editable": True,
                "sections": ["summary", "department_breakdown", "assumptions"],
                "unknown_value_policy": "mark_tbd",
                "untrusted_path": "../../secret",
            }
        },
    )
    spec = resolve_deliverable_spec(proposal)
    assert spec.output_format is DeliverableFormat.DOCX
    assert spec.audience == "producer"
    assert spec.layout == "department_tables"
    assert spec.sections == ["summary", "department_breakdown", "assumptions"]
    assert "untrusted_path" not in spec.model_dump()


def test_dynamic_deliverable_spec_rejects_layout_format_mismatch():
    proposal = issue_action_token(
        space_id="space_bad_spec",
        project_tag="general",
        user_id="dynamic_user",
        title="Invalid CSV slides",
        description="CSV cannot express one scene per slide",
        sources=[],
        action_type="create_scene_breakdown",
        output_format="csv",
        metadata={"deliverable_spec": {"layout": "one_scene_per_slide"}},
    )
    with pytest.raises(DeliverableSpecError, match="LAYOUT_NOT_SUPPORTED_FOR_FORMAT"):
        resolve_deliverable_spec(proposal)

@pytest.mark.parametrize(
    "act_type",
    [
        "create_call_sheet",
        "generate_shot_list",
        "stunt_risk_breakdown",
        "export_production_budget",
        "create_scene_breakdown",
        "generate_pitch_deck",
    ],
)
@pytest.mark.parametrize(
    ("output_format", "magic", "mime"),
    [
        ("pdf", b"%PDF", "application/pdf"),
        ("docx", b"PK", "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
        ("xlsx", b"PK", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
        ("pptx", b"PK", "application/vnd.openxmlformats-officedocument.presentationml.presentation"),
        ("csv", b"", "text/csv"),
        ("json", b"{", "application/json"),
    ],
)
def test_generate_complete_action_format_matrix(act_type, output_format, magic, mime):
    space_id = "space_5format_test"
    user = User(uid="user_tester", email="test@studiotower.ai", display_name="Lead Producer")

    title = f"Canonical {act_type} artifact"
    proposal = issue_action_token(
            space_id=space_id,
            project_tag="general",
            user_id=user.uid,
            title=title,
            description="Schema-first deterministic deliverable",
            sources=[],
            action_type=act_type,
            output_format=output_format,
            ttl_seconds=300,
        )

    exec_rec = ActionExecutionRecord(
            action_id=proposal.action_id,
            run_id=f"run_{proposal.action_id[:10]}",
            space_id=space_id,
            project_tag="general",
            user_id=user.uid,
            status=ActionExecutionStatus.PENDING,
        )
    store.create_action_execution_if_absent(exec_rec)

    run, updated_exec = DeliverableExecutionService.execute_action(proposal, user)
    assert run.status.value in ("completed", "awaiting_approval")
    assert len(run.output_artifact_ids) == 1

    art_id = run.output_artifact_ids[0]
    art = store.get_artifact(space_id, art_id)
    assert art is not None
    assert art.filename.endswith(f".{output_format}")
    assert art.media_type == mime

    blob = store.get_artifact_blob(space_id, art_id, art.filename)
    assert blob is not None
    if magic:
        assert blob.startswith(magic)


def test_export_production_budget_in_pdf():
    space_id = "space_budget_pdf_test"
    user = User(uid="user_tester_budget", email="producer@studiotower.ai", display_name="Line Producer")

    proposal = issue_action_token(
        space_id=space_id,
        project_tag="general",
        user_id=user.uid,
        title="Production Budget Estimate (PDF)",
        description="Department line items, equipment rates, and daily spend estimate in PDF format",
        sources=[],
        action_type="export_production_budget",
        output_format="pdf",
        ttl_seconds=300,
    )

    exec_rec = ActionExecutionRecord(
        action_id=proposal.action_id,
        run_id=f"run_{proposal.action_id[:10]}",
        space_id=space_id,
        project_tag="general",
        user_id=user.uid,
        status=ActionExecutionStatus.PENDING,
    )
    store.create_action_execution_if_absent(exec_rec)

    run, updated_exec = DeliverableExecutionService.execute_action(proposal, user)
    assert run.status.value == "completed"
    assert len(run.output_artifact_ids) == 1

    art_id = run.output_artifact_ids[0]
    art = store.get_artifact(space_id, art_id)
    assert art is not None
    assert art.filename.endswith(".pdf")
    assert art.media_type == "application/pdf"

    blob = store.get_artifact_blob(space_id, art_id, art.filename)
    assert blob is not None
    assert blob.startswith(b"%PDF")
    # Verify it is genuine budget content, not a call sheet
    assert b"CALL SHEET" not in blob
    assert b"EXT. EXPORT PRODUCTION BUDGET" not in blob


def test_space_lineage_normalization():
    from app.models.file_record import FileRecord, FileSourceType, IngestionStatus
    from app.models.space import Space
    from app.services.lineage_service import LineageService

    space_id = "space_lin_test"
    user = User(uid="lin_tester", email="lin@studiotower.ai", display_name="Auditor")

    space = Space(space_id=space_id, name="Lineage Space", owner_uid=user.uid, created_by=user.uid)
    store.create_space(space, user.uid)
    store.add_member(space_id, user.uid, role="owner")

    # Upload a source file
    source_f = FileRecord(
        file_id="file_src_1",
        space_id=space_id,
        filename="Test_Script.pdf",
        content_type="application/pdf",
        size_bytes=1024,
        sha256="abcdef123456",
        storage_path="path/to/src",
        project_tags=["general"],
        uploaded_by=user.uid,
        source_type=FileSourceType.USER_UPLOAD,
        upload_status="committed",
        ingestion_status=IngestionStatus.READY,
    )
    store.save_file(source_f)

    # Test lineage with tag="all"
    lin_all = LineageService.get_space_lineage(space_id, "all", user)
    assert len(lin_all["nodes"]) >= 1
    assert any(n["id"] == "file_src_1" for n in lin_all["nodes"])

    # Test lineage with tag="all-tracks"
    lin_tracks = LineageService.get_space_lineage(space_id, "all-tracks", user)
    assert len(lin_tracks["nodes"]) >= 1
    assert any(n["id"] == "file_src_1" for n in lin_tracks["nodes"])
